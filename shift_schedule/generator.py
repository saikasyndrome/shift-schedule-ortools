"""Config parse, heuristic engine, validation, Excel export."""

import yaml
import calendar
import copy
import datetime
import random
import re
import hashlib
import time
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def load_config(path="config/sample.yaml"):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)

class ShiftGenerator:
    def __init__(self, config):
        self.config   = config
        self.year     = config["schedule"]["year"]
        self.month    = config["schedule"]["month"]
        self.num_days = calendar.monthrange(self.year, self.month)[1]
        self.staff    = config["staff"]
        self.rules    = config["rules"]

        self.required_codes = [
            s["code"] for s in config["shift_types"] if s.get("required")
        ]

        self.hard_errors = []

        self.exceptions = {
            (e["name"], e["date"]): e
            for e in config.get("exceptions", [])
        }

        self.ojt_map = {}
        for t in config.get("ojt", []):
            for entry in t["staff"]:
                name  = entry["name"]
                shift = entry.get("shift", "")
                self.ojt_map[(t["date"], name)] = shift

        self.ojt_teacher_map = {}

        self.ojt_pairs = []

        self.ojt_partner = {}

        self.ojt_avoid_days = {}

        self.shift_map = {s["code"]: s for s in config["shift_types"]}

        self.result = {s["name"]: {} for s in self.staff}

        for entry in config.get("prev_month", []):
            name   = entry["name"]
            shifts = entry.get("last_shifts", [])[-5:]  # 最大5日だけ取る
            for i, code in enumerate(reversed(shifts)):
                self.result[name][-i] = code

        self.assign_warnings = []

    def date_str(self, day):
        return f"{self.year}-{self.month:02d}-{day:02d}"

    def is_exception(self, name, day):
        return self.exceptions.get((name, self.date_str(day)))

    def is_exception_applied(self, name, day):
        exc = self.is_exception(name, day)
        if not exc or exc["type"] == "夜勤不可":
            return False
        actual = self.base_code(self.result.get(name, {}).get(day, ""))
        return actual == exc["type"]

    def get_ojt_shift(self, name, day):
        return self.ojt_map.get((self.date_str(day), name), "")

    def is_ojt(self, name, day):
        return bool(self.get_ojt_shift(name, day))

    def get_ojt_assignment(self, name, day):
        date = self.date_str(day)
        if (date, name) in self.ojt_map:
            return self.ojt_map[(date, name)], "target"
        if (date, name) in self.ojt_teacher_map:
            return self.ojt_teacher_map[(date, name)], "teacher"
        return None, None

    def prev_shift(self, name, day):
        if day > 1:
            return self.result[name].get(day - 1, "")
        else:
            return self.result[name].get(0, "")  # 前月末日

    def consecutive_count(self, name, day):
        count = 0

        d = day - 1
        while d >= 1:
            s = self.result[name].get(d, "")
            base = self.base_code(s)
            if base in ("", "休み", "明け", "有給", "?"):
                return count
            count += 1
            d -= 1

        prev_key = 0
        while prev_key >= -4:
            s = self.result[name].get(prev_key, "")
            if s == "":
                break  # 前月データなし → カウント終了
            if s in ("休み", "明け", "有給", "?"):
                break
            count += 1
            prev_key -= 1

        return count

    def is_work_shift(self, code):
        base = self.base_code(code)
        return base != "休み" and base != ""

    def base_code(self, code):
        return code.replace("(OJT指導)", "").replace("(OJT)", "")

    def is_early(self, code):
        return self.shift_map.get(self.base_code(code), {}).get("type") == "early"

    def is_late(self, code):
        return self.shift_map.get(self.base_code(code), {}).get("type") == "late"

    def is_daytime(self, code):
        return self.base_code(code) in ("早1", "遅1", "早4", "遅4", "日勤")

    def floor_type_allowed(self, floor_list, s_floor, s_type):
        allowed = False
        for entry in floor_list:
            if isinstance(entry, dict):
                f, t = entry.get("floor"), entry.get("type")
            else:
                f, t = entry, None
            if f != s_floor:
                continue
            if t is None or t == s_type:
                allowed = True
        return allowed

    def floor_display(self, staff_info):
        floor_list = staff_info.get("floor", [])
        parts = []
        for entry in floor_list:
            if isinstance(entry, dict):
                f = entry.get("floor", "")
                t = entry.get("type")
                label = {"early": "早のみ", "late": "遅のみ"}.get(t, "")
                parts.append(f"{f}({label})" if label else f)
            else:
                parts.append(entry)

        ojt_target_raw = staff_info.get("ojt_target")
        if ojt_target_raw:
            for t_floor, t_type in self.parse_ojt_entries(ojt_target_raw):
                if self.floor_type_allowed(floor_list, t_floor, t_type or ""):
                    continue  # 既にfloorに含まれている(一人でも可)分は重複表示しない
                label = {"early": "早のみ", "late": "遅のみ"}.get(t_type, "")
                parts.append(f"{t_floor}({label})(OJT)" if label else f"{t_floor}(OJT)")

        return "・".join(parts)

    def _required_flexibility(self, staff_info, day):
        name = staff_info["name"]
        prev = self.prev_shift(name, day)
        avail = self.get_available_shifts(staff_info, day)
        if self.rules.get("no_late_to_early") and self.is_late(prev):
            avail = [s for s in avail if not (self.is_early(s) or self.base_code(s) == "日勤")]
        elif self.rules.get("no_late_to_early") and self.base_code(prev) == "日勤":
            avail = [s for s in avail if not self.is_early(s)]
        return len(set(avail) & set(self.required_codes))

    def get_available_shifts(self, staff_info, day):
        floor = staff_info["floor"]
        available = []
        for s in self.config["shift_types"]:
            code = s["code"]
            if code in ("明け", "休み"):
                continue
            if code == "夜勤":
                if not staff_info.get("night_ok", False):
                    continue
                available.append(code)
                continue
            if code == "日勤":
                if not staff_info.get("day_ok", False):
                    continue
                available.append(code)
                continue
            s_floor = s.get("floor", "")
            if s_floor and not self.floor_type_allowed(floor, s_floor, s.get("type", "")):
                continue
            available.append(code)

        matched_code = self.ojt_map.get((self.date_str(day), staff_info["name"]))
        if matched_code and matched_code not in available:
            available.append(matched_code)

        return available

    def count_nights(self, name):
        return sum(1 for s in self.result[name].values() if self.base_code(s) == "夜勤")

    def count_work_days(self, name):
        return sum(
            1 for k, v in self.result[name].items()
            if k >= 1 and self.is_work_shift(v)
        )

    def count_off_days(self, name):
        return sum(
            1 for k, v in self.result[name].items()
            if k >= 1 and v == "休み"
        )

    def future_forced_off_count(self, name, day):
        return sum(
            1 for d in range(day + 1, self.num_days + 1)
            if (exc := self.is_exception(name, d)) and exc["type"] == "休み"
        )

    def early_late_counts(self, name):
        early = late = 0
        for v in self.result[name].values():
            code = self.base_code(v)
            if code in ("早1", "早4"):
                early += 1
            elif code in ("遅1", "遅4"):
                late += 1
        return early, late

    def floor_counts(self, name):
        sakura = floor4 = 0
        for v in self.result[name].values():
            code = self.base_code(v)
            fl = self.shift_map.get(code, {}).get("floor")
            if fl == "さくら":
                sakura += 1
            elif fl == "4階":
                floor4 += 1
        return sakura, floor4

    def code_counts(self, name):
        counts = {"早1": 0, "遅1": 0, "早4": 0, "遅4": 0}
        for v in self.result[name].values():
            code = self.base_code(v)
            if code in counts:
                counts[code] += 1
        return counts

    def parse_ojt_value(self, value):
        if isinstance(value, dict):
            return value.get("floor"), value.get("type")
        return value, None

    def parse_ojt_entries(self, value):
        if isinstance(value, list):
            return [self.parse_ojt_value(v) for v in value]
        return [self.parse_ojt_value(value)]

    def ojt_floor_codes(self, floor, type_restrict=None):
        if floor == "夜勤":
            return [s["code"] for s in self.config["shift_types"] if s.get("type") == "night"]
        types = (type_restrict,) if type_restrict else ("early", "late")
        return [
            s["code"] for s in self.config["shift_types"]
            if s.get("floor") == floor and s.get("type") in types
        ]

    def _auto_assign_ojt(self):
        for target in self.staff:
            ojt_target_raw = target.get("ojt_target")
            if not ojt_target_raw:
                continue
            target_entries = self.parse_ojt_entries(ojt_target_raw)
            target_count = target.get("ojt_count", 0)
            unlimited = not target_count  # ojt_count未設定 = 目標なし、入れられる限り最大まで詰め込む

            target_name = target["name"]

            floor_codes = {}
            for ojt_floor, target_type in target_entries:
                codes = self.ojt_floor_codes(ojt_floor, target_type)
                if codes:
                    floor_codes[ojt_floor] = (codes, target_type)
            if not floor_codes:
                self.assign_warnings.append(
                    f"⚠ {target_name}: ojt_target: {ojt_target_raw} に該当する早番/遅番シフトコードが"
                    f"shift_typesに見つからず、OJT自動マッチング不可"
                )
                continue

            already = sum(1 for (_, name) in self.ojt_map if name == target_name)
            if unlimited:
                remaining = self.num_days  # 事実上無制限(月の日数以上は入り得ない)
            else:
                remaining = target_count - already
                if remaining <= 0:
                    continue

            teachers = []
            for t in self.staff:
                if t["name"] == target_name:
                    continue
                t_raw = t.get("ojt_teacher")
                if not t_raw:
                    continue
                for t_floor, t_type in self.parse_ojt_entries(t_raw):
                    if t_floor not in floor_codes:
                        continue
                    codes, target_type = floor_codes[t_floor]
                    pair_codes = [c for c in codes if t_type is None or self.shift_map.get(c, {}).get("type") == t_type]
                    if pair_codes:
                        teachers.append((t, pair_codes))

            if not teachers:
                if not unlimited:
                    self.assign_warnings.append(
                        f"⚠ {target_name}: OJT {target_count}回 目標のうち {already}回のみ確保 "
                        f"(ojt_teacher: {ojt_target_raw} に対応する先生が見つからず、またはtype不一致で "
                        f"{remaining}回 自動マッチング不可)"
                    )
                continue

            half = max(1, self.num_days // 2)
            step = max(1, half // (remaining + 1))
            day_order = list(range(step, half + 1, step))
            day_order += [d for d in range(1, self.num_days + 1) if d not in day_order]

            filled = 0

            def _try_fill(day, respect_avoid):
                nonlocal filled
                date = self.date_str(day)
                if (date, target_name) in self.ojt_map or (date, target_name) in self.ojt_teacher_map:
                    return False
                t_exc = self.is_exception(target_name, day)
                if t_exc and t_exc["type"] == "休み":
                    return False
                if day == 1 and self.result[target_name].get(0) in ("夜勤", "明け"):
                    return False
                if respect_avoid and day in self.ojt_avoid_days.get(target_name, ()):
                    return False

                for teacher, pair_codes in teachers:
                    teacher_name = teacher["name"]
                    if (date, teacher_name) in self.ojt_map or (date, teacher_name) in self.ojt_teacher_map:
                        continue
                    te_exc = self.is_exception(teacher_name, day)
                    if te_exc and te_exc["type"] == "休み":
                        continue
                    if day == 1 and self.result[teacher_name].get(0) in ("夜勤", "明け"):
                        continue
                    if respect_avoid and day in self.ojt_avoid_days.get(teacher_name, ()):
                        continue

                    e_cnt, l_cnt = self.early_late_counts(target_name)
                    early_codes = [c for c in pair_codes if self.is_early(c)]
                    late_codes  = [c for c in pair_codes if self.is_late(c)]
                    if e_cnt <= l_cnt and early_codes:
                        shift_code = early_codes[0]
                    elif late_codes:
                        shift_code = late_codes[0]
                    else:
                        shift_code = pair_codes[0]

                    self.ojt_map[(date, target_name)] = shift_code
                    self.ojt_teacher_map[(date, teacher_name)] = shift_code
                    self.ojt_pairs.append((date, target_name, teacher_name, shift_code))
                    self.ojt_partner[(date, target_name)] = teacher_name
                    self.ojt_partner[(date, teacher_name)] = target_name
                    filled += 1
                    return True
                return False

            for day in day_order:
                if filled >= remaining:
                    break
                _try_fill(day, respect_avoid=True)

            if filled < remaining:
                for day in day_order:
                    if filled >= remaining:
                        break
                    _try_fill(day, respect_avoid=False)

            if not unlimited and filled < remaining:
                self.assign_warnings.append(
                    f"⚠ {target_name}: OJT {target_count}回 目標のうち {already + filled}回のみ確保 "
                    f"(先生・日程マッチング不可で {remaining - filled}回 不足)"
                )

    def generate(self, seed=42, resume_from_day=1, prefix_result=None):
        random.seed(seed)
        self._seed = seed  # 平準化オフセット等、seedごとに違うパターンを試したい箇所で使う
        self.assign_warnings = []
        self._check_work_off_bounds_sane()
        if prefix_result:
            for name, days in prefix_result.items():
                self.result[name].update(days)
        self._auto_assign_ojt()

        night_required = self.rules.get("night_required_per_day", 1)

        for day in range(resume_from_day, self.num_days + 1):
            night_count_today = 0
            required_filled = {code: False for code in self.required_codes}

            order = list(range(len(self.staff)))
            random.shuffle(order)

            pending = []  # (idx, prev, exc, remaining_days)
            for idx in order:
                staff_info = self.staff[idx]
                name       = staff_info["name"]
                prev = self.prev_shift(name, day)
                exc = self.is_exception(name, day)

                def _override_ojt_warning(reason):
                    ojt_shift, ojt_role = self.get_ojt_assignment(name, day)
                    if ojt_shift:
                        role_label = "対象者" if ojt_role == "target" else "先生役"
                        self.assign_warnings.append(
                            f"⚠ {name} {day}日: OJT({role_label})予定でしたが{reason}のため振替が必要です"
                        )

                if exc and exc["type"] in ("休み", "有給"):
                    self.result[name][day] = exc["type"]
                    _override_ojt_warning(f"{exc['type']}例外優先")
                    continue
                if exc and exc["type"] in self.shift_map and exc["type"] not in ("明け", "休み"):
                    forced_code = exc["type"]
                    if self._structurally_capable(staff_info, forced_code):
                        self.result[name][day] = forced_code
                        if forced_code == "夜勤": night_count_today += 1
                        if forced_code in required_filled: required_filled[forced_code] = True
                        _override_ojt_warning(f"例外({forced_code}指定)")
                        continue
                    else:
                        self.assign_warnings.append(
                            f"⚠ {name} {day}日: 例外({forced_code}指定)予定でしたが適性"
                            f"(floor/night_ok/day_ok)を満たさないため無視されました"
                        )

                if self.base_code(prev) == "夜勤":
                    self.result[name][day] = "明け"
                    _override_ojt_warning("夜勤明け優先")
                    continue

                if prev == "明け":
                    self.result[name][day] = "休み"
                    _override_ojt_warning("明け休み優先")
                    continue

                remaining_days = self.num_days - day + 1

                max_consecutive = staff_info.get("max_consecutive", self.rules["max_consecutive"])
                namehash = sum(ord(c) for c in name)
                stagger = (namehash * 7 + self._seed * 13) % min(3, max_consecutive)
                effective_limit = max(1, max_consecutive - stagger)
                if self.consecutive_count(name, day) >= effective_limit:
                    self.result[name][day] = "休み"
                    _override_ojt_warning("最大連続勤務日数制限")
                    continue

                pending.append((idx, prev, exc, remaining_days))

            prepared = []  # (idx, prev, exc, remaining_days, available)
            for idx, prev, exc, remaining_days in pending:
                staff_info = self.staff[idx]
                name       = staff_info["name"]

                available = self.get_available_shifts(staff_info, day)

                if not (exc and exc["type"] == "日勤"):
                    available = [s for s in available if s != "日勤"]

                if exc and exc["type"] == "夜勤不可":
                    available = [s for s in available if s != "夜勤"]

                prev_base = self.base_code(prev)
                if self.rules.get("no_late_to_early") and self.is_late(prev):
                    available = [s for s in available if not (self.is_early(s) or self.base_code(s) == "日勤")]
                elif self.rules.get("no_late_to_early") and prev_base == "日勤":
                    available = [s for s in available if not self.is_early(s)]

                if self.count_nights(name) >= staff_info.get("max_nights", 99):
                    available = [s for s in available if s != "夜勤"]

                if day < self.num_days:
                    off_needed_after_today = self.rules["off_days"] - self.count_off_days(name)
                    if off_needed_after_today >= remaining_days - 1:
                        available = [s for s in available if s != "夜勤"]

                next_exc = self.is_exception(name, day + 1) if day + 1 <= self.num_days else None
                if next_exc and next_exc["type"] != "夜勤不可":
                    available = [s for s in available if s != "夜勤"]
                else:
                    next2_exc = self.is_exception(name, day + 2) if day + 2 <= self.num_days else None
                    if next2_exc and next2_exc["type"] not in ("休み", "有給", "夜勤不可"):
                        available = [s for s in available if s != "夜勤"]

                prepared.append((idx, prev, exc, remaining_days, available))

            def _has_other_plan(idx, day, exc):
                name = self.staff[idx]["name"]
                ojt_shift, _ = self.get_ojt_assignment(name, day)
                return bool(ojt_shift)

            def _off_days_room(idx, day):
                name = self.staff[idx]["name"]
                return self.rules["off_days"] - self.count_off_days(name) - self.future_forced_off_count(name, day)

            def _would_exceed_off_budget(idx, day):
                name = self.staff[idx]["name"]
                day_plus2 = day + 2
                if day_plus2 <= self.num_days:
                    day2_exc = self.is_exception(name, day_plus2)
                    chain_off_counts = not (day2_exc and day2_exc["type"] == "有給")
                else:
                    chain_off_counts = False  # 明け→休みが翌月に跨るため今月のoff_daysには影響しない
                predicted = (
                    self.count_off_days(name)
                    + self.future_forced_off_count(name, day)
                    + (1 if chain_off_counts else 0)
                )
                return predicted > self.rules["off_days"]

            night_order = sorted(
                prepared,
                key=lambda t: (
                    1 if _has_other_plan(t[0], day, t[2]) else 0,
                    0 if self._required_flexibility(self.staff[t[0]], day) == 0 else 1,
                    1 if _would_exceed_off_budget(t[0], day) else 0,
                    self.count_nights(self.staff[t[0]]["name"]),
                    self.staff[t[0]].get("leader", False),
                    -_off_days_room(t[0], day),
                ),
            )
            after_night = []  # (idx, prev, exc, remaining_days, available) 夜勤確定に回らなかった人
            for idx, prev, exc, remaining_days, available in night_order:
                if night_count_today < night_required and "夜勤" in available:
                    name = self.staff[idx]["name"]
                    self.result[name][day] = "夜勤"
                    night_count_today += 1
                    if _has_other_plan(idx, day, exc):
                        self._warn_if_forced_shift_exception_overridden(name, day, exc, "夜勤最優先確保")
                        ojt_shift, ojt_role = self.get_ojt_assignment(name, day)
                        if ojt_shift:
                            role_label = "対象者" if ojt_role == "target" else "先生役"
                            self.assign_warnings.append(
                                f"⚠ {name} {day}日: OJT({role_label})予定でしたが夜勤最優先確保のため振替が必要です"
                            )
                else:
                    after_night.append((idx, prev, exc, remaining_days, available))

            date_str_today = self.date_str(day)
            avail_by_name = {
                self.staff[i]["name"]: avail
                for i, prev0, exc0, remaining0, avail in after_night
            }

            still_pending = []  # (idx, prev, available)
            for idx, prev, exc, remaining_days, available in after_night:
                staff_info = self.staff[idx]
                name       = staff_info["name"]

                ojt_shift, ojt_role = self.get_ojt_assignment(name, day)
                if ojt_shift:
                    partner_name = self.ojt_partner.get((date_str_today, name))
                    partner_ok = True
                    if partner_name is not None:
                        partner_avail = avail_by_name.get(partner_name)
                        partner_ok = partner_avail is not None and ojt_shift in partner_avail
                    if ojt_shift in available and partner_ok:
                        suffix = "(OJT)" if ojt_role == "target" else "(OJT指導)"
                        self.result[name][day] = f"{ojt_shift}{suffix}"
                        if ojt_shift == "夜勤": night_count_today += 1
                        if ojt_shift in required_filled: required_filled[ojt_shift] = True
                        continue
                    else:
                        role_label = "対象者" if ojt_role == "target" else "先生役"
                        if ojt_shift not in available:
                            reason = f"{ojt_shift}が割当不可"
                        else:
                            reason = "相方(先生役/対象者)が同日に確保できない"
                        self.assign_warnings.append(
                            f"⚠ {name} {day}日: OJT({role_label})予定でしたがルール制約"
                            f"({reason})のため振替が必要です"
                        )
                        if not self.can_assign_code(staff_info, ojt_shift):
                            available = [s for s in available if s != ojt_shift]

                still_pending.append((idx, prev, available))

            for i in range(len(still_pending)):
                idx, prev, available = still_pending[i]
                still_pending[i] = (idx, prev, [s for s in available if s != "夜勤"])

            remaining = list(still_pending)
            while True:
                needed_codes = [c for c in self.required_codes if not required_filled.get(c, True)]
                if not needed_codes:
                    break
                code_candidates = {}
                for c in needed_codes:
                    positions = [p for p, (idx, prev, avail) in enumerate(remaining) if c in avail]
                    if positions:
                        code_candidates[c] = positions
                if not code_candidates:
                    break  # 誰も入れないコードは、後段の必須シフト未充足チェックに任せる
                target_code = min(code_candidates, key=lambda c: len(code_candidates[c]))
                positions = code_candidates[target_code]

                def _balance_score(p):
                    idx_p, prev_p, avail_p = remaining[p]
                    name = self.staff[idx_p]["name"]
                    counts = self.code_counts(name)
                    personal_codes = [c for c in avail_p if c in self.required_codes]
                    own_max = max((counts.get(c, 0) for c in personal_codes), default=0)
                    return own_max - counts.get(target_code, 0)

                positions.sort(key=lambda p: (
                    0 if self._required_flexibility(self.staff[remaining[p][0]], day) == 1 else 1,
                    -_balance_score(p),
                    self._required_flexibility(self.staff[remaining[p][0]], day),
                    self.staff[remaining[p][0]].get("leader", False),
                ))
                pos = positions[0]
                idx, prev, avail = remaining[pos]
                name = self.staff[idx]["name"]
                self.result[name][day] = target_code
                required_filled[target_code] = True
                del remaining[pos]

            random.shuffle(remaining)
            
            for idx, prev, available in remaining:
                staff_info = self.staff[idx]
                name       = staff_info["name"]
                candidates = available

                if not candidates:
                    self.result[name][day] = "?"
                    reason = self._unassignable_reason(staff_info, day, prev, night_count_today >= night_required)
                    msg = f"⚠ {name} {day}日: 配置不可 [{reason}]"
                    self.assign_warnings.append(msg)
                    self.hard_errors.append(msg)
                    continue

                if staff_info.get("leader"):
                    leaders = [s["name"] for s in self.staff if s.get("leader") and s["name"] != name]
                    other_daytime = any(self.is_daytime(self.result[ln].get(day, "")) for ln in leaders)
                    if not other_daytime:
                        day_candidates = [s for s in candidates if self.is_daytime(s)]
                        if day_candidates:
                            candidates = day_candidates

                BIAS = 4.0
                counts = self.code_counts(name)
                eligible_counts = [counts[self.base_code(c)] for c in candidates if self.base_code(c) in counts]
                max_count = max(eligible_counts) if eligible_counts else 0

                weights = []
                for c in candidates:
                    bc = self.base_code(c)
                    if bc in counts:
                        deficit = max_count - counts[bc]
                        weights.append(BIAS ** deficit)
                    else:
                        weights.append(1.0)

                chosen = random.choices(candidates, weights=weights, k=1)[0] if len(candidates) > 1 else candidates[0]
                self.result[name][day] = chosen

                if chosen in required_filled: required_filled[chosen] = True

            for code in self.required_codes:
                if not required_filled.get(code, True):
                    detail = self._required_shift_gap_detail(day, code)
                    msg = f"⚠ {day}日: {code} 0名 (最低1名 未達) - {detail}"
                    self.assign_warnings.append(msg)
                    self.hard_errors.append(msg)

        self._backfill_off_days()
        return self.result

    def _backfill_off_days(self):
        CONVERTIBLE = ("早1", "遅1", "早4", "遅4", "日勤")
        OFF_CAP = self.rules.get("max_consecutive_off", 2)

        def rest_count_that_day(day):
            return sum(
                1 for s in self.staff
                if self.result[s["name"]].get(day, "") == "休み"
            )

        def _off_like(code):
            return self.base_code(code) in ("休み", "有給")

        def _window_has_forced_exception(name, day):
            for d in range(day - OFF_CAP, day + OFF_CAP + 1):
                exc = self.is_exception(name, d)
                if exc and exc["type"] in ("休み", "有給"):
                    return True
            return False

        def _would_exceed_off_cap(name, day):
            if _window_has_forced_exception(name, day):
                return False
            for start in range(day - OFF_CAP, day + 1):
                end = start + OFF_CAP
                count = 0
                for d in range(start, end + 1):
                    if d == day:
                        count += 1
                    elif 1 <= d <= self.num_days and _off_like(self.result[name].get(d, "")):
                        count += 1
                if count > OFF_CAP:
                    return True
            return False

        order = list(range(len(self.staff)))
        random.shuffle(order)  # 誰から優先的に良い(空いている)日を選べるかをseedごとに変える
        for idx in order:
            name = self.staff[idx]["name"]
            deficit = self.rules["off_days"] - self.count_off_days(name)
            if deficit <= 0:
                continue

            candidates = []
            for day in range(1, self.num_days + 1):
                code = self.result[name].get(day, "")
                base = self.base_code(code)
                if base not in CONVERTIBLE:
                    continue
                if code != base:
                    continue  # OJT中の日は相方が一人になるため変換しない
                exc = self.is_exception(name, day)
                if exc and exc["type"] not in ("休み", "有給"):
                    continue  # 手動の勤務強制指定はそのまま尊重する
                if base in self.required_codes:
                    others_cover = any(
                        o["name"] != name
                        and self.base_code(self.result[o["name"]].get(day, "")) == base
                        for o in self.staff
                    )
                    if not others_cover:
                        continue  # その日の唯一の担当者→変換すると必須シフトが欠員になる
                candidates.append(day)

            candidates.sort(key=rest_count_that_day)
            for day in candidates:
                if deficit <= 0:
                    break
                if _would_exceed_off_cap(name, day):
                    continue  # 連続休みがmax_consecutive_offを超えてしまう日はスキップ
                self.result[name][day] = "休み"
                deficit -= 1

    def _warn_if_forced_shift_exception_overridden(self, name, day, exc, reason):
        if exc and exc["type"] in self.shift_map:
            self.assign_warnings.append(
                f"⚠ {name} {day}日: 例外({exc['type']}指定)予定でしたが{reason}のため振替が必要です"
            )

    def _unassignable_reason(self, staff_info, day, prev, night_assigned):
        reasons = []
        if self.is_late(prev):
            reasons.append("遅番翌日")
        if self.count_nights(staff_info["name"]) >= staff_info.get("max_nights", 99):
            reasons.append("夜勤MAX超過")
        if night_assigned:
            reasons.append("夜勤割当済み")
        return "・".join(reasons) if reasons else "候補なし"

    def _check_work_off_bounds_sane(self):
        off_days = self.rules["off_days"]
        if off_days > self.num_days:
            raise ValueError(
                f"rules不整合: off_days({off_days}) が総日数({self.num_days})を超えています"
            )

    def yukyu_total(self, name):
        return sum(
            1 for (n, _d), e in self.exceptions.items()
            if n == name and e["type"] == "有給"
        )

    def _structurally_capable(self, staff_info, code):
        if code == "夜勤":
            return staff_info.get("night_ok", False)
        if code == "日勤":
            return staff_info.get("day_ok", False)
        s_def = self.shift_map.get(code, {})
        s_floor = s_def.get("floor", "")
        if not s_floor:
            return True
        return self.floor_type_allowed(staff_info.get("floor", []), s_floor, s_def.get("type", ""))

    def can_assign_code(self, staff_info, code):
        s_def = self.shift_map.get(code, {})
        s_floor = s_def.get("floor", "")
        if not s_floor:
            return True  # floor制約が無いコード(日勤等)は誰でも対象になり得る
        return self.floor_type_allowed(staff_info.get("floor", []), s_floor, s_def.get("type", ""))

    def effective_work_target(self, name):
        return self.num_days - self.rules["off_days"]

    def _required_shift_gap_detail(self, day, code):
        s_def   = self.shift_map.get(code, {})
        s_floor = s_def.get("floor", "")
        s_type  = s_def.get("type", "")
        eligible = [
            s for s in self.staff
            if self.floor_type_allowed(s.get("floor", []), s_floor, s_type)
        ]
        if not eligible:
            return f"{code}を担当できるスタッフが0人です (floor設定を確認してください)"

        parts = [f"{s['name']}:{self.result[s['name']].get(day, '') or '?'}" for s in eligible]
        return f"担当可能{len(eligible)}名 [" + ", ".join(parts) + "]"

    def _night_gap_detail(self, day):
        night_eligible = [s for s in self.staff if s.get("night_ok", False)]
        if not night_eligible:
            return "夜勤可能なスタッフが0人です (night_ok: true のスタッフを追加してください)"

        parts = []
        for s in night_eligible:
            name  = s["name"]
            code  = self.result[name].get(day, "") or "?"
            base  = self.base_code(code)
            max_n = s.get("max_nights", 99)

            if base == "明け":
                reason = "明け(前日夜勤の休養中)"
            elif base == "休み":
                reason = "休み"
            elif base == "有給":
                reason = "有給"
            elif code.endswith("(OJT)") or code.endswith("(OJT指導)"):
                reason = code
            elif self.count_nights(name) >= max_n:
                reason = f"{code}(夜勤MAX{max_n}回消化済み)"
            else:
                reason = code
            parts.append(f"{name}:{reason}")

        return "夜勤可能" + str(len(night_eligible)) + "名 [" + ", ".join(parts) + "]"

    def validate(self):
        warnings = list(self.assign_warnings)
        hard = list(self.hard_errors)  # generate()内で検出済みの「配置不可」を引き継ぐ

        for s in self.staff:
            name = s["name"]
            for day in range(1, self.num_days + 1):
                shift = self.result[name].get(day, "")
                prev  = self.result[name].get(day - 1, "") if day > 1 else ""

                if self.is_late(prev) and self.is_early(shift):
                    msg = f"⚠ {name} {day}日: 遅番→早番NG ({prev}→{shift})"
                    warnings.append(msg)
                    hard.append(msg)
                if shift == "?":
                    warnings.append(f"⚠ {name} {day}日: 未配置 (手動確認必要)")

            night_count = self.count_nights(name)
            max_n = s.get("max_nights", 99)
            if night_count > max_n:
                msg = f"⚠ {name}: 夜勤{night_count}回 (MAX{max_n}回超過)"
                warnings.append(msg)
                hard.append(msg)

            off_count = self.count_off_days(name)
            off_target = self.rules["off_days"]
            if off_count < off_target:
                msg = f"⚠ {name}: 休み{off_count}日 (目標{off_target}日に{off_target - off_count}日不足)"
                warnings.append(msg)
                hard.append(msg)
            elif off_count > off_target:
                msg = f"⚠ {name}: 休み{off_count}日 (目標{off_target}日より{off_count - off_target}日超過)"
                warnings.append(msg)
                hard.append(msg)

            work_count  = self.count_work_days(name)
            work_target = self.effective_work_target(name)
            if work_count > work_target:
                msg = f"⚠ {name}: 勤務{work_count}日 (目標{work_target}日より{work_count - work_target}日超過)"
                warnings.append(msg)
                hard.append(msg)
            elif work_count < work_target:
                yukyu_note = f"、うち有給{self.yukyu_total(name)}日" if self.yukyu_total(name) else ""
                msg = f"⚠ {name}: 勤務{work_count}日 (目標{work_target}日に{work_target - work_count}日不足{yukyu_note})"
                warnings.append(msg)
                hard.append(msg)

        leaders = [s["name"] for s in self.staff if s.get("leader")]
        for day in range(1, self.num_days + 1):
            if not any(self.is_daytime(self.result[ln].get(day, "")) for ln in leaders):
                warnings.append(f"⚠ {day}日: リーダー日中不在")

        night_required = self.rules.get("night_required_per_day", 1)
        for day in range(1, self.num_days + 1):
            night_count = sum(1 for s in self.staff if self.base_code(self.result[s["name"]].get(day, "")) == "夜勤")
            if night_count < night_required:
                detail = self._night_gap_detail(day)
                msg = f"⚠ {day}日: 夜勤 {night_count}名 (最低{night_required}名 未達) - {detail}"
                warnings.append(msg)
                hard.append(msg)

        for date, target_name, teacher_name, shift_code in self.ojt_pairs:
            day = int(date.split("-")[-1])
            target_actual  = self.result[target_name].get(day, "")
            teacher_actual = self.result[teacher_name].get(day, "")
            if target_actual == f"{shift_code}(OJT)" and teacher_actual != f"{shift_code}(OJT指導)":
                warnings.append(
                    f"⚠ {target_name} {day}日: OJT予定でしたが担当の先生({teacher_name})が"
                    f"出勤不可のため指導者なしで実施されました"
                )

        def dedup(items):
            seen = set()
            unique = []
            for w in items:
                if w not in seen:
                    seen.add(w)
                    unique.append(w)
            return unique

        self.hard_errors = dedup(hard)
        return dedup(warnings)

    def summary(self):
        rows = []
        for s in self.staff:
            name   = s["name"]
            shifts = self.result[name]
            current = {k: v for k, v in shifts.items() if k >= 1}
            work       = sum(1 for v in current.values() if self.is_work_shift(v))
            nikkin     = sum(1 for v in current.values() if self.base_code(v) == "日勤")
            nights     = sum(1 for v in current.values() if self.base_code(v) == "夜勤")
            ake        = sum(1 for v in current.values() if v == "明け")
            off        = sum(1 for v in current.values() if v == "休み")
            yukyu      = sum(1 for v in current.values() if v == "有給")
            ojt        = sum(1 for v in current.values() if v.endswith("(OJT)"))
            ojt_teach  = sum(1 for v in current.values() if v.endswith("(OJT指導)"))
            unassigned = sum(1 for v in current.values() if v == "?")
            hours      = sum(self.shift_map.get(self.base_code(v), {}).get("hours", 0) for v in current.values())
            early1     = sum(1 for v in current.values() if self.base_code(v) == "早1")
            late1      = sum(1 for v in current.values() if self.base_code(v) == "遅1")
            early4     = sum(1 for v in current.values() if self.base_code(v) == "早4")
            late4      = sum(1 for v in current.values() if self.base_code(v) == "遅4")
            if not self.can_assign_code(s, "早1") and early1 == 0: early1 = "x"
            if not self.can_assign_code(s, "遅1") and late1  == 0: late1  = "x"
            if not self.can_assign_code(s, "早4") and early4 == 0: early4 = "x"
            if not self.can_assign_code(s, "遅4") and late4  == 0: late4  = "x"
            if not s.get("night_ok", False):
                nights = "x"
                ake    = "x"
            rows.append({
                "name": name, "work": work, "nikkin": nikkin, "nights": nights,
                "ake": ake, "off": off, "yukyu": yukyu, "hours": hours,
                "ojt": ojt, "ojt_teach": ojt_teach, "unassigned": unassigned,
                "early1": early1, "late1": late1, "early4": early4, "late4": late4,
            })
        return rows

class ExcelExporter:
    C_NAVY     = "1B2A3B"
    C_BLUE_HDR = "1A5276"
    C_WHITE    = "FFFFFF"
    C_GRAY_LT  = "F2F3F4"
    C_INPUT    = "EBF5FB"
    C_GREEN_LT = "D5F5E3"
    C_OJT      = "F8C471"  # OJT対象者・OJT指導(先生役)は組んだペア同士が一目で分かるよう同一色にする

    SHIFT_COLORS = {
        "早1":  "D5F5E3",
        "遅1":  "D6EAF8",
        "早4":  "FEF9E7",
        "遅4":  "FDEBD0",
        "日勤": "F0B27A",
        "夜勤": "F9EBEA",
        "明け": "E8DAEF",
        "休み": "F2F3F4",
        "有給": "D1F2EB",
        "?":    "F1948A",
    }
    DOW_JP = ["月", "火", "水", "木", "金", "土", "日"]

    def __init__(self, generator):
        self.gen      = generator
        self.year     = generator.year
        self.month    = generator.month
        self.num_days = generator.num_days
        self.staff    = generator.staff
        self.result   = generator.result
        self.wb       = Workbook()

    def fl(self, hex_c):
        return PatternFill("solid", fgColor=hex_c)

    def fn(self, size=10, bold=False, color="000000", italic=False):
        return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)

    def al(self, h="center", v="center", wrap=True):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    def border(self):
        s = Side(style="thin", color="BFBFBF")
        return Border(left=s, right=s, top=s, bottom=s)

    def sc(self, ws, row, col, value="", bg=None, fg="000000",
           size=10, bold=False, align="center", italic=False):
        c = ws.cell(row=row, column=col, value=value)
        if bg:
            c.fill = self.fl(bg)
        c.font      = self.fn(size=size, bold=bold, color=fg, italic=italic)
        c.alignment = self.al(h=align)
        c.border    = self.border()
        return c

    def build_shift_sheet(self):
        ws = self.wb.active
        ws.title = f"Shift_{self.year}{self.month:02d}"
        ws.sheet_view.showGridLines = False

        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 8
        for d in range(1, self.num_days + 1):
            ws.column_dimensions[get_column_letter(d + 3)].width = 5.2
        sum_labels = ["出勤日数", "総時間(h)", "早1", "遅1", "早4", "遅4", "日勤", "夜勤", "明け", "休み", "有給", "OJT", "OJT指導", "未配置", "状態"]
        for i in range(len(sum_labels)):
            ws.column_dimensions[get_column_letter(self.num_days + 4 + i)].width = 9

        last = get_column_letter(self.num_days + 3 + len(sum_labels))
        ws.merge_cells(f"B1:{last}1")
        c = ws["B1"]
        c.value     = f"📅  シフト表 — {self.year}年{self.month}月"
        c.font      = self.fn(size=13, bold=True, color=self.C_WHITE)
        c.fill      = self.fl(self.C_NAVY)
        c.alignment = self.al()
        ws.row_dimensions[1].height = 30

        ws.row_dimensions[2].height = 20
        self.sc(ws, 2, 2, "氏名", bg=self.C_NAVY, fg=self.C_WHITE, bold=True, size=9)
        self.sc(ws, 2, 3, "担当", bg=self.C_NAVY, fg=self.C_WHITE, bold=True, size=9)
        for d in range(1, self.num_days + 1):
            dt = datetime.date(self.year, self.month, d)
            bg = "8B0000" if dt.weekday() >= 5 else self.C_NAVY
            self.sc(ws, 2, d + 3, d, bg=bg, fg=self.C_WHITE, bold=True, size=8)
        for i, lbl in enumerate(sum_labels):
            self.sc(ws, 2, self.num_days + 4 + i, lbl,
                    bg=self.C_BLUE_HDR, fg=self.C_WHITE, bold=True, size=8)

        ws.row_dimensions[3].height = 14
        self.sc(ws, 3, 2, "", bg=self.C_NAVY)
        self.sc(ws, 3, 3, "", bg=self.C_NAVY)
        for d in range(1, self.num_days + 1):
            dt  = datetime.date(self.year, self.month, d)
            dow = self.DOW_JP[dt.weekday()]
            bg  = "8B0000" if dt.weekday() >= 5 else self.C_NAVY
            self.sc(ws, 3, d + 3, dow, bg=bg, fg=self.C_WHITE, size=7)
        for i in range(len(sum_labels)):
            self.sc(ws, 3, self.num_days + 4 + i, "", bg=self.C_BLUE_HDR)

        data_start = 4
        for si, s in enumerate(self.staff):
            name  = s["name"]
            floor = self.gen.floor_display(s)
            row   = data_start + si
            ws.row_dimensions[row].height = 20

            self.sc(ws, row, 2, name,  bg="ECF0F1", bold=True, align="left")
            self.sc(ws, row, 3, floor, bg="ECF0F1", size=8)

            for d in range(1, self.num_days + 1):
                code    = self.result[name].get(d, "")
                is_ojt_target  = code.endswith("(OJT)")
                is_ojt_teacher = code.endswith("(OJT指導)")
                if is_ojt_teacher or is_ojt_target:
                    bg = self.C_OJT
                else:
                    bg = self.SHIFT_COLORS.get(code, self.C_WHITE)
                dt = datetime.date(self.year, self.month, d)
                if not code and dt.weekday() >= 5:
                    bg = "FAEBD7"
                fg = "C0392B" if code == "?" else "2C3E50"
                display_code = f"{code}(希)" if code and self.gen.is_exception_applied(name, d) else code
                self.sc(ws, row, d + 3, display_code, bg=bg, fg=fg, bold=bool(code), size=7)

            current = {k: v for k, v in self.result[name].items() if k >= 1}
            work       = sum(1 for v in current.values() if self.gen.is_work_shift(v))
            early1     = sum(1 for v in current.values() if self.gen.base_code(v) == "早1")
            late1      = sum(1 for v in current.values() if self.gen.base_code(v) == "遅1")
            early4     = sum(1 for v in current.values() if self.gen.base_code(v) == "早4")
            late4      = sum(1 for v in current.values() if self.gen.base_code(v) == "遅4")
            if not self.gen.can_assign_code(s, "早1") and early1 == 0: early1 = "x"
            if not self.gen.can_assign_code(s, "遅1") and late1  == 0: late1  = "x"
            if not self.gen.can_assign_code(s, "早4") and early4 == 0: early4 = "x"
            if not self.gen.can_assign_code(s, "遅4") and late4  == 0: late4  = "x"
            nikkin     = sum(1 for v in current.values() if self.gen.base_code(v) == "日勤")
            nights     = sum(1 for v in current.values() if self.gen.base_code(v) == "夜勤")
            ake        = sum(1 for v in current.values() if v == "明け")
            off        = sum(1 for v in current.values() if v == "休み")
            yukyu      = sum(1 for v in current.values() if v == "有給")
            ojt        = sum(1 for v in current.values() if v.endswith("(OJT)"))
            ojt_teach  = sum(1 for v in current.values() if v.endswith("(OJT指導)"))
            unassigned = sum(1 for v in current.values() if v == "?")
            hours      = sum(self.gen.shift_map.get(self.gen.base_code(v), {}).get("hours", 0) for v in current.values())
            max_n      = s.get("max_nights", 99)
            off_target = self.gen.rules["off_days"]

            status_parts = []
            if nights > max_n:
                status_parts.append("夜勤超過")
            if off != off_target:
                status_parts.append(f"休み不一致({off}/{off_target})")
            if unassigned > 0:
                status_parts.append(f"未配置{unassigned}件")
            status = "⚠ " + "・".join(status_parts) if status_parts else "✅ OK"

            self.sc(ws, row, self.num_days + 4,  work,       bg=self.C_INPUT,    bold=True)
            self.sc(ws, row, self.num_days + 5,  hours,      bg=self.C_GREEN_LT, bold=True)
            self.sc(ws, row, self.num_days + 6,  early1,     bg="D5F5E3",        bold=True)
            self.sc(ws, row, self.num_days + 7,  late1,      bg="D6EAF8",        bold=True)
            self.sc(ws, row, self.num_days + 8,  early4,     bg="FEF9E7",        bold=True)
            self.sc(ws, row, self.num_days + 9,  late4,      bg="FDEBD0",        bold=True)
            self.sc(ws, row, self.num_days + 10, nikkin,     bg="F0B27A",        bold=True)
            self.sc(ws, row, self.num_days + 11, nights,     bg="F9EBEA",        bold=True)
            self.sc(ws, row, self.num_days + 12, ake,        bg="E8DAEF",        bold=True)
            self.sc(ws, row, self.num_days + 13, off,
                    bg="F1948A" if off != off_target else "F2F3F4", bold=True)
            self.sc(ws, row, self.num_days + 14, yukyu,      bg="D1F2EB",        bold=True)
            self.sc(ws, row, self.num_days + 15, ojt,        bg="F9E79F",        bold=True)
            self.sc(ws, row, self.num_days + 16, ojt_teach,  bg="AED6F1",        bold=True)
            self.sc(ws, row, self.num_days + 17, unassigned,
                    bg="F1948A" if unassigned else self.C_GRAY_LT, bold=True)
            self.sc(ws, row, self.num_days + 18, status,     bg=self.C_WHITE,    bold=True)

        hc_row = data_start + len(self.staff)
        ws.row_dimensions[hc_row].height = 18
        ws.merge_cells(f"B{hc_row}:C{hc_row}")
        c = ws[f"B{hc_row}"]
        c.value     = "出勤人数合計"
        c.font      = self.fn(size=8, bold=True, color=self.C_WHITE)
        c.fill      = self.fl(self.C_NAVY)
        c.alignment = self.al()
        c.border    = self.border()

        for d in range(1, self.num_days + 1):
            count = sum(
                1 for s in self.staff
                if self.result[s["name"]].get(d, "") not in ("", "休み", "明け", "有給", "?")
            )
            self.sc(ws, hc_row, d + 3, count, bg=self.C_NAVY, fg=self.C_WHITE, bold=True, size=9)
        for i in range(len(sum_labels)):
            self.sc(ws, hc_row, self.num_days + 4 + i, "", bg=self.C_NAVY)

        leg_row = hc_row + 2
        ws.row_dimensions[leg_row].height = 17
        c = ws.cell(row=leg_row, column=2, value="凡例:")
        c.font      = self.fn(size=8, bold=True)
        c.alignment = self.al(h="left")
        legend = [
            (3,  "早1=さくら早番", "D5F5E3"),
            (5,  "遅1=さくら遅番", "D6EAF8"),
            (7,  "早4=4階早番",    "FEF9E7"),
            (9,  "遅4=4階遅番",    "FDEBD0"),
            (11, "日勤",           "F0B27A"),
            (13, "夜勤",           "F9EBEA"),
            (15, "明け",           "E8DAEF"),
            (17, "休み",           "F2F3F4"),
            (19, "有給",           "D1F2EB"),
            (21, "XX(OJT)=OJT対象者",   self.C_OJT),
            (23, "XX(OJT指導)=OJT先生役", self.C_OJT),
            (25, "XX(希)=exceptions指定日", "F2F3F4"),
            (27, "? = 未配置",     "F1948A"),
        ]
        for col, label, color in legend:
            self.sc(ws, leg_row, col, label, bg=color, fg="2C3E50", bold=True, size=8)

        return ws

    def build_summary_sheet(self):
        ws = self.wb.create_sheet("Summary")
        ws.sheet_view.showGridLines = False

        for ci, w in enumerate([2, 18, 10, 12, 9, 9, 9, 9, 10, 10, 12, 10, 10, 8, 8, 8, 8, 12], start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w

        ws.merge_cells("B1:R1")
        c = ws["B1"]
        c.value     = f"📊  月次サマリー — {self.year}年{self.month}月"
        c.font      = self.fn(size=13, bold=True, color=self.C_WHITE)
        c.fill      = self.fl(self.C_BLUE_HDR)
        c.alignment = self.al()
        ws.row_dimensions[1].height = 30

        hdrs = ["氏名", "担当", "出勤日数", "総時間(h)", "早1", "遅1", "早4", "遅4",
                "日勤回数", "夜勤回数", "明け回数", "休み日数", "有給日数",
                "OJT回数", "OJT指導回数", "未配置", "状態"]
        ws.row_dimensions[3].height = 22
        for i, h in enumerate(hdrs):
            self.sc(ws, 3, i + 2, h, bg=self.C_BLUE_HDR, fg=self.C_WHITE, bold=True, size=9)

        summary = self.gen.summary()
        for si, row_data in enumerate(summary):
            row     = 4 + si
            s       = self.staff[si]
            floor   = self.gen.floor_display(s)
            max_n      = s.get("max_nights", 99)
            off_target = self.gen.rules["off_days"]
            status_parts = []
            if isinstance(row_data["nights"], int) and row_data["nights"] > max_n:
                status_parts.append("夜勤超過")
            if row_data["off"] != off_target:
                status_parts.append(f"休み不一致({row_data['off']}/{off_target})")
            if row_data["unassigned"] > 0:
                status_parts.append(f"未配置{row_data['unassigned']}件")
            status = "⚠ " + "・".join(status_parts) if status_parts else "✅ OK"
            ws.row_dimensions[row].height = 20

            self.sc(ws, row, 2,  row_data["name"],    bg=self.C_INPUT,    bold=True, align="left")
            self.sc(ws, row, 3,  floor,                bg=self.C_INPUT,    size=9)
            self.sc(ws, row, 4,  row_data["work"],     bg=self.C_INPUT,    bold=True)
            self.sc(ws, row, 5,  row_data["hours"],    bg=self.C_GREEN_LT, bold=True)
            self.sc(ws, row, 6,  row_data["early1"],   bg="D5F5E3",        bold=True)
            self.sc(ws, row, 7,  row_data["late1"],    bg="D6EAF8",        bold=True)
            self.sc(ws, row, 8,  row_data["early4"],   bg="FEF9E7",        bold=True)
            self.sc(ws, row, 9,  row_data["late4"],    bg="FDEBD0",        bold=True)
            self.sc(ws, row, 10, row_data["nikkin"],   bg="F0B27A",        bold=True)
            self.sc(ws, row, 11, row_data["nights"],   bg="F9EBEA",        bold=True)
            self.sc(ws, row, 12, row_data["ake"],      bg="E8DAEF",        bold=True)
            self.sc(ws, row, 13, row_data["off"],
                    bg="F1948A" if row_data["off"] != off_target else "F2F3F4", bold=True)
            self.sc(ws, row, 14, row_data["yukyu"],     bg="D1F2EB",        bold=True)
            self.sc(ws, row, 15, row_data["ojt"],       bg="F9E79F",        bold=True)
            self.sc(ws, row, 16, row_data["ojt_teach"], bg="AED6F1",        bold=True)
            self.sc(ws, row, 17, row_data["unassigned"],
                    bg="F1948A" if row_data["unassigned"] else self.C_GRAY_LT, bold=True)
            self.sc(ws, row, 18, status,               bg=self.C_WHITE,    bold=True)

        tot = 4 + len(self.staff)
        ws.row_dimensions[tot].height = 20
        ws.merge_cells(f"B{tot}:C{tot}")
        self.sc(ws, tot, 2, "合計", bg=self.C_NAVY, fg=self.C_WHITE, bold=True)
        for i, key in enumerate(["work", "hours", "early1", "late1", "early4", "late4",
                                  "nikkin", "nights", "ake", "off", "yukyu",
                                  "ojt", "ojt_teach", "unassigned"]):
            self.sc(ws, tot, 4 + i, sum(r[key] for r in summary if isinstance(r[key], int)),
                    bg=self.C_NAVY, fg=self.C_WHITE, bold=True)
        self.sc(ws, tot, 18, "", bg=self.C_NAVY)

        return ws

    def build_warnings_sheet(self, warnings):
        ws = self.wb.create_sheet("Warnings")
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 65

        ws.merge_cells("B1:B1")
        c = ws["B1"]
        c.value     = "⚠️  ルール違反・未配置チェック結果"
        c.font      = self.fn(size=12, bold=True, color=self.C_WHITE)
        c.fill      = self.fl("922B21")
        c.alignment = self.al(h="left")
        ws.row_dimensions[1].height = 28

        if not warnings:
            c = ws.cell(row=3, column=2, value="✅ ルール違反・未配置なし")
            c.font      = self.fn(size=11, bold=True, color="1E8449")
            c.alignment = self.al(h="left")
        else:
            for i, w in enumerate(warnings):
                row = 3 + i
                ws.row_dimensions[row].height = 18
                is_unassigned = "配置不可" in w or "未配置" in w
                bg = "F1948A" if is_unassigned else "FDEDEC"
                fg = "7B241C" if is_unassigned else "922B21"
                c = ws.cell(row=row, column=2, value=w)
                c.font      = self.fn(size=10, bold=is_unassigned, color=fg)
                c.fill      = self.fl(bg)
                c.alignment = self.al(h="left")
                c.border    = self.border()

        return ws

    def export(self, output_path):
        self.build_shift_sheet()
        self.build_summary_sheet()
        warnings = self.gen.validate()
        self.build_warnings_sheet(warnings)
        self.wb.save(output_path)
        return warnings

def _hard_error_days(hard_errors):
    days = set()
    for msg in hard_errors:
        m = re.search(r"(\d+)日[:：]", msg)
        if m:
            days.add(int(m.group(1)))
    return days

def _warnings_before_day(warnings, day_limit):
    kept = []
    for msg in warnings:
        m = re.search(r"(\d+)日[:：]", msg)
        if m and int(m.group(1)) < day_limit:
            kept.append(msg)
    return kept

def _balance_penalty(gen):
    codes = ("早1", "遅1", "早4", "遅4")
    penalty = 0
    for s in gen.staff:
        counts = gen.code_counts(s["name"])
        capable = [c for c in codes if gen.can_assign_code(s, c)]
        if len(capable) < 2:
            continue
        vals = [counts[c] for c in capable]
        avg = sum(vals) / len(vals)
        penalty += sum((v - avg) ** 2 for v in vals)
    return penalty

def _ojt_shortfall(gen):
    total = 0
    for s in gen.staff:
        target_count = s.get("ojt_count", 0)
        if not (s.get("ojt_target") and target_count):
            continue
        name = s["name"]
        achieved = sum(
            1 for d in range(1, gen.num_days + 1)
            if str(gen.result[name].get(d, "")).endswith("(OJT)")
        )
        total += max(0, target_count - achieved)
    return total

def _derived_seed(base_seed, counter):
    h = hashlib.sha256(f"{base_seed}:{counter}".encode()).hexdigest()
    return int(h[:16], 16)

def _predict_night_busy_days(config, seed):
    stripped = copy.deepcopy(config)
    for s in stripped.get("staff", []):
        s.pop("ojt_target", None)
        s.pop("ojt_teacher", None)
        s.pop("ojt_count", None)
    stripped.pop("ojt", None)

    probe = ShiftGenerator(stripped)
    try:
        probe.generate(seed=seed)
    except Exception:
        return {}

    busy = {}
    for s in config.get("staff", []):
        if not s.get("night_ok"):
            continue
        name = s["name"]
        days = {
            d for d in range(1, probe.num_days + 1)
            if probe.base_code(probe.result[name].get(d)) in ("夜勤", "明け")
        }
        if days:
            busy[name] = days
    return busy

def generate_best(config, base_seed=1, max_attempts=1000, day_attempts=30, balance_seconds=0):
    BALANCE_SEARCH_SECONDS = balance_seconds
    full_tries = 0
    seed_counter = 0
    balance_deadline = None

    uses_ojt = (
        any(s.get("ojt_target") for s in config.get("staff", []))
        and any(s.get("ojt_teacher") for s in config.get("staff", []))
    )
    ojt_avoid_days = _predict_night_busy_days(config, base_seed) if uses_ojt else {}

    def new_gen():
        g = ShiftGenerator(config)
        g.ojt_avoid_days = ojt_avoid_days
        return g

    def alloc_seed():
        nonlocal seed_counter
        seed_counter += 1
        return _derived_seed(base_seed, seed_counter)

    FUTILITY_PATIENCE = 20

    def day_backtrack(gen, budget):
        used = 0
        while used < budget:
            problem_days = _hard_error_days(gen.hard_errors)
            first_bad_day = min(problem_days) if problem_days else None
            if not first_bad_day or first_bad_day <= 1:
                break  # 月次集計系のみが残っている、または1日目からダメ→日単位では直せない
            prefix = {
                s["name"]: {d: v for d, v in gen.result[s["name"]].items() if 1 <= d < first_bad_day}
                for s in gen.staff
            }
            prefix_warnings = _warnings_before_day(gen.assign_warnings, first_bad_day)
            improved = False
            no_improve_streak = 0
            while used < budget:
                used += 1
                cand = new_gen()
                cand.generate(seed=alloc_seed(), resume_from_day=first_bad_day, prefix_result=prefix)
                cand.assign_warnings = prefix_warnings + cand.assign_warnings
                cand.validate()
                if not cand.hard_errors:
                    return cand, True
                if len(cand.hard_errors) < len(gen.hard_errors):
                    gen = cand
                    improved = True
                    break  # 直ったこの日を確定させ、次の(新しい)問題日を探しにいく
                no_improve_streak += 1
                if no_improve_streak >= FUTILITY_PATIENCE:
                    break  # 直近FUTILITY_PATIENCE回改善ゼロ→この日は固定されていて直せないと判断
            if not improved:
                break  # この問題日はbudget内では直せなかった→この系統は諦める
        return gen, False

    best_gen, best_count = None, None
    best_success, best_metric = None, None

    FULL_FUTILITY_PATIENCE = min(max_attempts, 300)
    no_improve_full_streak = 0

    def consider_success(cand):
        nonlocal best_success, best_metric
        metric = (_ojt_shortfall(cand), _balance_penalty(cand))
        if best_success is None or metric < best_metric:
            best_success, best_metric = cand, metric
        return metric == (0, 0)

    while True:
        if best_success is not None:
            if balance_deadline is None:
                balance_deadline = time.time() + BALANCE_SEARCH_SECONDS
            if time.time() >= balance_deadline:
                break
        elif full_tries >= max_attempts:
            break
        elif no_improve_full_streak >= FULL_FUTILITY_PATIENCE:
            break  # 直近FULL_FUTILITY_PATIENCE回のフル生成で改善ゼロ→構造的に無理と判断し早期終了

        full_tries += 1
        gen = new_gen()
        gen.generate(seed=base_seed if full_tries == 1 else alloc_seed())
        gen.validate()

        if not gen.hard_errors:
            if consider_success(gen):
                return best_success, True
            no_improve_full_streak = 0
        else:
            gen, done = day_backtrack(gen, day_attempts)
            if done:
                if consider_success(gen):
                    return best_success, True
                no_improve_full_streak = 0
            elif best_success is None:
                if best_count is None or len(gen.hard_errors) < best_count:
                    best_gen, best_count = gen, len(gen.hard_errors)
                    no_improve_full_streak = 0
                else:
                    no_improve_full_streak += 1

    if best_success is not None:
        return best_success, True
    return best_gen, False
